"""
excel_export.py
Feeds an Excel workbook that other label-printing software (BarTender,
NiceLabel, etc.) can point at as a data source - no manual copy/paste.

Two ways rows get in:
  1. export_pullout_rows() - dumps every line item of a pull-out record
     straight from the database (used by the "Export to Label Excel"
     button).
  2. append_scanned_payload() - takes a raw scanned barcode string (same
     format printed on the Lot List / QA Acceptance Lot PDF) and appends
     it as a row (used by the Label Scan Station in the Packing List tab).

Both write into the same workbook/sheet so everything ends up in one
place, in the same column layout, for downstream tools to consume.

Separately, export_lot_list_backup() and export_pullout_log_backup()
dump the full Lot List / pull-out history tables as plain backup
spreadsheets - unrelated to the label-data format above, just a
record-keeping export the user can keep alongside the .db file backup.
"""
import os
from openpyxl import Workbook, load_workbook
from . import label_payload as lp
from . import database as db

SHEET_NAME = "Label Data"
HEADERS = [lp.LINE_FIELD_LABELS[k] for k in lp.LINE_FIELD_ORDER]


def _open_or_create(xlsx_path):
    if os.path.exists(xlsx_path):
        wb = load_workbook(xlsx_path)
        if SHEET_NAME not in wb.sheetnames:
            ws = wb.create_sheet(SHEET_NAME)
            ws.append(HEADERS)
        else:
            ws = wb[SHEET_NAME]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(HEADERS)
    return wb, ws


def append_scanned_payload(xlsx_path, payload):
    """payload: the raw pipe-delimited string scanned off a printed barcode."""
    data = lp.parse_line_payload(payload)
    wb, ws = _open_or_create(xlsx_path)
    ws.append([data.get(k, "") for k in lp.LINE_FIELD_ORDER])
    wb.save(xlsx_path)
    return data


def export_pullout_rows(xlsx_path, pullout, pullout_lots):
    """Writes one row per source lot line item for this pull-out, without
    requiring a physical scan - useful for back-office bulk export."""
    wb, ws = _open_or_create(xlsx_path)
    count = 0
    for pl in pullout_lots:
        payload = lp.build_line_payload(pullout, pl)
        data = lp.parse_line_payload(payload)
        ws.append([data.get(k, "") for k in lp.LINE_FIELD_ORDER])
        count += 1
    wb.save(xlsx_path)
    return count


LOT_LIST_BACKUP_HEADERS = [
    "Customer", "Part No", "Part Name", "Rev", "Heat No", "M/C No", "Lot No",
    "Lot Qty", "Date of OQC", "RTV", "Balance Lot Qty", "Lot Date", "Latest Pull-Out Date",
]


def export_lot_list_backup(xlsx_path):
    """Full, unfiltered dump of every lot in the Lot List - a plain backup
    export, independent of the label-data format used elsewhere."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lot List"
    ws.append(LOT_LIST_BACKUP_HEADERS)
    rows = db.list_lots()
    for r in rows:
        ws.append([
            r["customer_name"], r["part_no"], r["part_name"] or "", r["rev"] or "",
            r["heat_no"] or "", r["mc_no"] or "", r["route_card_lot_no"] or "",
            r["input_lot_qty"], r["date_of_oqc"] or "", "RTV" if r["is_rtv"] else "",
            r["balance_lot_qty"], r["scan_date"], r["last_pullout_date"] or "",
        ])
    wb.save(xlsx_path)
    return len(rows)


PULLOUT_LOG_BACKUP_HEADERS = [
    "Status", "Customer", "Part No (Customer)", "Rev", "Packaging Lot No", "Heat No",
    "Source Lot Numbers", "Packaging Qty", "Packaging Date", "PO Number", "Prepared By",
]


def export_pullout_log_backup(xlsx_path):
    """Full, unfiltered dump of the entire pull-out / Packing List history -
    a plain backup export, independent of the label-data format used
    elsewhere."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pullout Log"
    ws.append(PULLOUT_LOG_BACKUP_HEADERS)
    rows = db.list_pullouts()
    for p in rows:
        source_lots = db.get_pullout_lots(p["id"])
        source_str = ", ".join(sl["route_card_lot_no"] or "" for sl in source_lots)
        heat_no = p["heat_no"] or (source_lots[-1]["heat_no"] if source_lots else "")
        display_part_no = p["customer_part_no"] or p["part_no"]
        ws.append([
            p["status"] or "OPEN", p["customer_name"], display_part_no, p["rev"] or "",
            p["packaging_lot_no"], heat_no or "", source_str, p["packaging_qty"],
            p["packaging_date"], p["po_number"], p["prepared_by"],
        ])
    wb.save(xlsx_path)
    return len(rows)
