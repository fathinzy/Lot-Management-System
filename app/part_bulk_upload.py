"""
part_bulk_upload.py
Mass-registers parts from an Excel file. A row is only imported if its
Customer is already in the Customer Register AND its Code is already in
that customer's Code Register - rows that fail either check are skipped
and reported back with a reason, everything else is imported.
"""
from openpyxl import Workbook, load_workbook
from . import database as db

TEMPLATE_HEADERS = [
    "Customer", "Part Number", "Customer Part Number", "Part Name", "Rev",
    "Code", "Material Supplier", "Material Type", "Default Lot Qty", "Lot No Rule",
]

REQUIRED_HEADERS = ["customer", "part number", "code"]


def create_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts"
    ws.append(TEMPLATE_HEADERS)
    ws.append(["Globex Inc", "PN-1000", "CUST-PN-1000", "Bracket Assembly", "A",
               "AC-01", "SupplierX", "Al6061", 100, "Standard CSR"])
    wb.save(path)


def _cell(row, colidx, name, default=None):
    idx = colidx.get(name)
    if idx is None or idx >= len(row):
        return default
    value = row[idx]
    return value if value is not None else default


def _to_str(value):
    return str(value).strip() if value is not None else ""


def bulk_upload_parts(filepath):
    """Returns (uploaded, skipped):
    uploaded: list of (row_number, part_no)
    skipped:  list of (row_number, part_no, reason)
    Raises ValueError if the file is missing required columns.
    """
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    colidx = {
        _to_str(h).lower(): i for i, h in enumerate(header_row) if h is not None
    }
    missing = [h for h in REQUIRED_HEADERS if h not in colidx]
    if missing:
        raise ValueError(
            "Excel file is missing required column(s): " + ", ".join(missing) +
            ". Use the Download Template button for the expected layout."
        )

    uploaded = []
    skipped = []

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None for v in row):
            continue  # blank row

        customer_name = _to_str(_cell(row, colidx, "customer"))
        part_no = _to_str(_cell(row, colidx, "part number"))
        code_val = _to_str(_cell(row, colidx, "code"))

        if not customer_name or not part_no:
            skipped.append((row_num, part_no, "Missing Customer or Part Number"))
            continue

        customer = db.get_customer_by_name(customer_name)
        if not customer:
            skipped.append((row_num, part_no, f"Customer '{customer_name}' not registered"))
            continue

        if not code_val:
            skipped.append((row_num, part_no, "Code is required (must already be in the "
                                                "Code Register for this customer)"))
            continue

        codes = db.list_codes(customer["id"])
        code_row = next((c for c in codes if c["code"] == code_val), None)
        if not code_row:
            skipped.append((row_num, part_no,
                             f"Code '{code_val}' not registered for '{customer_name}'"))
            continue

        customer_part_no = _to_str(_cell(row, colidx, "customer part number")) or None
        part_name = _to_str(_cell(row, colidx, "part name")) or None
        rev = _to_str(_cell(row, colidx, "rev"))
        material_supplier = _to_str(_cell(row, colidx, "material supplier")) or None
        material_type = _to_str(_cell(row, colidx, "material type")) or None

        default_qty_raw = _cell(row, colidx, "default lot qty")
        try:
            default_qty = float(default_qty_raw) if default_qty_raw not in (None, "") else None
        except (ValueError, TypeError):
            default_qty = None

        rule_name = _to_str(_cell(row, colidx, "lot no rule"))
        rule_id = None
        if rule_name:
            rules = db.list_rules(customer["id"])
            rule_row = next((r for r in rules if r["rule_name"] == rule_name), None)
            rule_id = rule_row["id"] if rule_row else None

        db.upsert_part(
            customer_id=customer["id"], part_no=part_no, part_name=part_name, rev=rev,
            code_id=code_row["id"], material_supplier=material_supplier,
            material_type=material_type, default_lot_qty=default_qty,
            default_rule_id=rule_id, customer_part_no=customer_part_no,
        )
        uploaded.append((row_num, part_no))

    return uploaded, skipped
